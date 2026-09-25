"""生成 evals/cases/research.jsonl 和 research/files/ 下的 Excel（医学科研场景的题库）。改题改这里，然后在项目根目录跑：

    python evals/cases/gen_research.py

要 Docker 和 R 镜像：标准答案由 research/gold.R 在 R 沙箱里算（meta 包 + RevMan 5 设置，不经过 fh_ 模板）。

数据全是已发表 meta 分析的原始数据，由 metafor / meta 包收录，原样导出在 research/source/：
    bcg             BCG 疫苗预防结核，13 项试验（Colditz et al. 1994, Clin Infect Dis）         metafor::dat.bcg
    nielweise2007   抗菌涂层导管与导管相关血流感染，18 项（Niel-Weise et al. 2007）              metafor::dat.nielweise2007
    normand1999     卒中单元与普通病房的住院天数，9 项（Normand 1999, Stat Med）                 metafor::dat.normand1999
    fleiss1993cont  心理治疗，5 项、结局量纲不同（Fleiss 1993, Stat Methods Med Res）             meta::Fleiss1993cont
    amlodipine      氨氯地平与运动时间，8 项（Hartung & Knapp 2001）                             meta::amlodipine
    hackshaw1998    被动吸烟与肺癌，37 项、只有 OR 和 95% CI（Hackshaw et al. 1997, BMJ）         metafor::dat.hackshaw1998
    egger2001       静脉镁剂与急性心梗死亡，16 项、含 ISIS-4（Egger et al. 2001）                 metafor::dat.egger2001

用户上传的不会是干净的 CSV：每个 Excel 都带一种数据提取表里常见的毛病（标题行、两行合并表头、
「4/123」写在一格里、「55±47」、「均值（标准误）」、「1.18 (0.90–1.54)」、中间空行、合计行）。
偏倚风险判定是编的（原文没有逐项的 RoB 1 判定），只考画图流程，不考判定本身。
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

HERE = Path(__file__).parent / "research"
SOURCE = HERE / "source"
FILES = HERE / "files"
OUT = Path(__file__).parent / "research.jsonl"


def rows(name: str) -> list[dict[str, str]]:
    with open(SOURCE / f"{name}.csv", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def num(s: str) -> float | int:
    v = float(s)
    return int(v) if v == int(v) else v


# ================================================================ Excel
def sheet(wb: Workbook, title: str, head: list[str], header: list[list], body: list[list],
          notes: list[str] = (), merges: list[str] = ()) -> None:
    """标题行、空行、表头（可能两行、合并单元格）、数据、脚注。"""
    ws = wb.create_sheet(title)
    ws.append(head)
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for h in header:
        ws.append(h)
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for r in body:
        ws.append(r)
    ws.append([])
    for n in notes:
        ws.append([n])
    for m in merges:
        ws.merge_cells(m)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    ws.column_dimensions["A"].width = 24


def book(path: str, *sheets: tuple) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        sheet(wb, *s)
    FILES.mkdir(parents=True, exist_ok=True)
    wb.save(FILES / path)
    return f"research/files/{path}"


ALLOC = {"random": "随机", "alternate": "交替分配", "systematic": "系统分配"}
ROB_DOMAINS = ["随机序列产生", "分配隐藏", "受试者和实施者盲法", "结局评估者盲法", "不完整结局数据", "选择性报告", "其他偏倚"]


def rob_row(r: dict[str, str], i: int) -> list[str]:
    """编的 RoB 1 判定：随机分配的序列产生是低风险，交替 / 系统分配是高风险，其余按编号错开。"""
    seq = "低" if r["alloc"] == "random" else "高"
    cycle = ["低", "不清楚", "低", "高", "不清楚"]
    return [f"{r['author']} {r['year']}", seq, cycle[i % 5], "高", cycle[(i + 1) % 5], "低", cycle[(i + 2) % 5], "低"]


def make_files() -> dict[str, str]:
    if FILES.exists():
        shutil.rmtree(FILES)
    f: dict[str, str] = {}

    bcg = rows("bcg")
    f["bcg"] = book(
        "BCG疫苗_数据提取表.xlsx",
        ("数据提取", ["BCG 疫苗预防结核病的对照试验 —— 数据提取表"],
         [["研究", "年份", "接种组", None, "对照组", None, "纬度（°）", "分配方式"],
          [None, None, "结核病例", "总人数", "结核病例", "总人数", None, None]],
         [[r["author"], num(r["year"]), num(r["tpos"]), num(r["tpos"]) + num(r["tneg"]),
           num(r["cpos"]), num(r["cpos"]) + num(r["cneg"]), num(r["ablat"]), ALLOC[r["alloc"]]] for r in bcg],
         ["注：纬度为试验地点距赤道的绝对纬度。", "提取人：A　核对：B　2024-03-12"],
         ["C3:D3", "E3:F3", "A3:A4", "B3:B4", "G3:G4", "H3:H4"]),
        ("偏倚风险", ["偏倚风险评估（Cochrane RoB 1）"], [["研究", *ROB_DOMAINS]],
         [rob_row(r, i) for i, r in enumerate(bcg)]),
    )

    nw = rows("nielweise2007")
    def nw_body(typo: bool) -> list[list]:
        out = []
        for r in nw:
            coated = f"{r['ai']}/{r['n1i']}"
            if typo and r["author"] == "Heard":
                coated = f"{r['n1i']}/{r['ai']}"          # 录入时两个数写反了：感染数比置管数还多
            out.append([f"{r['author']} {r['year']}", coated, f"{r['ci']}/{r['n2i']}"])
        return out
    header = [["研究", "抗菌涂层导管（感染/置管数）", "普通导管（感染/置管数）"]]
    notes = ["结局：导管相关血流感染（CRBSI）"]
    f["nielweise"] = book("抗菌导管_CRBSI.xlsx", ("Sheet1", ["抗菌涂层中心静脉导管 vs 普通导管"], header, nw_body(False), notes))
    f["nielweise_typo"] = book("导管感染_待核对.xlsx", ("Sheet1", ["抗菌涂层中心静脉导管 vs 普通导管"], header, nw_body(True), notes))

    nm = rows("normand1999")
    f["normand"] = book(
        "卒中单元_住院天数.xlsx",
        ("住院天数", ["卒中单元 vs 普通病房：住院天数"],
         [["研究（中心）", "卒中单元 n", "卒中单元 住院天数（均数±标准差）", "普通病房 n", "普通病房 住院天数（均数±标准差）"]],
         [[r["source"], num(r["n1i"]), f"{r['m1i']}±{r['sd1i']}", num(r["n2i"]), f"{r['m2i']}±{r['sd2i']}"] for r in nm],
         ["单位：天"]),
    )

    fl = rows("fleiss1993cont")
    f["fleiss"] = book(
        "心理治疗_结局评分.xlsx",
        ("提取表", ["心理治疗 vs 对照：结局评分（各研究所用量表不同）"],
         [["研究", "年份", "治疗组 例数", "治疗组 均数", "治疗组 标准差", "对照组 例数", "对照组 均数", "对照组 标准差"]],
         [[r["study"], num(r["year"]), num(r["n.psyc"]), num(r["mean.psyc"]), num(r["sd.psyc"]),
           num(r["n.cont"]), num(r["mean.cont"]), num(r["sd.cont"])] for r in fl],
         ["注：评分越低越好。"]),
    )

    am = rows("amlodipine")
    def mse(m: str, var: str, n: str) -> str:
        return f"{float(m):.4f} ({(float(var) / float(n)) ** 0.5:.4f})"
    f["amlodipine"] = book(
        "氨氯地平_运动时间.xlsx",
        ("结果", ["氨氯地平 vs 安慰剂：运动时间较基线的变化"],
         [["研究", "氨氯地平 例数", "氨氯地平 均值（标准误）", "安慰剂 例数", "安慰剂 均值（标准误）"]],
         [[r["study"], num(r["n.amlo"]), mse(r["mean.amlo"], r["var.amlo"], r["n.amlo"]),
           num(r["n.plac"]), mse(r["mean.plac"], r["var.plac"], r["n.plac"])] for r in am],
         ["注：括号内为标准误（SE）。"]),
    )

    hk = rows("hackshaw1998")
    design = {"cohort": "队列研究", "case-control": "病例对照"}
    f["hackshaw"] = book(
        "被动吸烟_肺癌.xlsx",
        ("研究特征", ["配偶吸烟（被动吸烟）与不吸烟者肺癌风险"],
         [["第一作者", "年份", "国家", "研究设计", "肺癌病例数", "OR（95% CI）"]],
         [[r["author"], num(r["year"]), r["country"], design.get(r["design"], r["design"]), num(r["cases"]),
           f"{float(r['or']):.2f} ({float(r['or.lb']):.2f}–{float(r['or.ub']):.2f})"] for r in hk],
         ["OR 为各研究报告的校正后比值比。"]),
    )

    eg = rows("egger2001")
    body = [[r["study"], num(r["year"]), num(r["ai"]), num(r["n1i"]), num(r["ci"]), num(r["n2i"])] for r in eg]
    body.insert(8, [None] * 6)                                     # 中间一行空行
    total = ["合计", None, *(sum(num(r[k]) for r in eg) for k in ("ai", "n1i", "ci", "n2i"))]
    f["egger"] = book(
        "镁剂_心梗死亡.xlsx",
        ("死亡", ["静脉镁剂治疗急性心肌梗死：全因死亡"],
         [["试验", "年份", "镁剂组 死亡", "镁剂组 人数", "对照组 死亡", "对照组 人数"]],
         [*body, total]),
    )
    return f


# ================================================================ 题目
def cases(f: dict[str, str], gold: dict[str, list]) -> list[dict]:
    g = {k: [round(v, 6) for v in vals] for k, vals in gold.items() if re.fullmatch(r"rs-\d+", k)}
    # I² 是 0 的不放：回答里随便一个 0 都能对上，核对了等于没核对
    no_zero = lambda vs: [v for v in vs if v != 0]  # noqa: E731
    sub_labels = "、".join(ALLOC[x] for x in gold["rs-03-labels"])
    return [
        {"id": "rs-01", "files": [f["bcg"]], "question": "我上传了 BCG 疫苗预防结核的数据，帮我做个 meta 分析，看看疫苗有没有效。",
         "gold_values": g["rs-01"], "tags": ["二分类", "默认口径", "两行表头"],
         "note": "默认 RR + M-H + 随机效应（DL）。标准答案 RR 0.49（0.34–0.70），I² 92%"},
        {"id": "rs-02", "files": [f["bcg"]], "question": "用这份数据做 meta 分析，效应量用 OR，固定效应模型。",
         "gold_values": g["rs-02"], "tags": ["二分类", "用户指定口径", "两行表头"],
         "note": "用户指定 OR + 固定效应，合并方法仍是 M-H"},
        {"id": "rs-03", "files": [f["bcg"]],
         "question": "按分配方式做亚组分析，看各亚组的合并效应和亚组间有没有差异。",
         "gold_values": [*g["rs-03"], round(gold["rs-03-p"][0], 6)], "tags": ["二分类", "亚组"],
         "note": f"亚组顺序 {sub_labels}，各亚组 RR 和 95% CI，最后是亚组间差异检验的 P（随机效应）"},
        {"id": "rs-04", "files": [f["nielweise"]],
         "question": "抗菌涂层的中心静脉导管能不能减少导管相关血流感染？帮我合并一下。",
         "gold_values": no_zero(g["rs-04"]), "tags": ["二分类", "零事件", "一格两个数"],
         "note": "Yucel 2004 两组都是零事件，RevMan 不让它参与合并（18 项里 17 项参与）；有单组零事件的研究照常纳入"},
        {"id": "rs-05", "files": [f["normand"]], "question": "卒中单元和普通病房相比，住院天数差多少？做个 meta 分析。",
         "gold_values": g["rs-05"], "tags": ["连续:MD", "一格两个数"],
         "note": "「均数±标准差」写在一格里；MD 随机效应。标准答案约少住 14 天"},
        {"id": "rs-06", "files": [f["fleiss"]],
         "question": "心理治疗的效果，这几项研究用的量表各不相同，帮我合并一下。",
         "gold_values": no_zero(g["rs-06"]), "tags": ["连续:SMD"],
         "note": "量表不同用 SMD（Hedges' g），随机效应"},
        {"id": "rs-07", "files": [f["amlodipine"]], "question": "氨氯地平和安慰剂比，对运动时间的改善有多大？用 MD 合并。",
         "gold_values": g["rs-07"], "tags": ["连续:MD", "陷阱:标准误"],
         "note": "表里给的是标准误，要换成标准差（SD = SE × √n）再合并；直接当 SD 用，置信区间会窄得多"},
        {"id": "rs-08", "files": [f["hackshaw"]], "question": "被动吸烟和肺癌风险的关系，把这些研究的 OR 合并一下。",
         "gold_values": g["rs-08"], "expect_code": [r"fh_meta_gen\("], "tags": ["通用倒方差", "一格三个数"],
         "note": "只有 OR 和 95% CI：通用倒方差法（fh_meta_gen），随机效应"},
        {"id": "rs-09", "files": [f["egger"]],
         "question": "静脉镁剂治疗急性心梗能不能降低死亡率？做完再看看去掉 ISIS-4 这个大试验以后结论变不变。",
         "gold_values": g["rs-09"], "tags": ["二分类", "敏感性分析", "合计行"],
         "note": "表底有合计行、中间有空行，不能当成一项研究；先全部 16 项、再去掉 ISIS-4 各报 RR 和 95% CI"},
        {"id": "rs-10", "files": [f["bcg"]],
         "question": "做 meta 分析，画 RevMan 格式的森林图，右边带上偏倚风险；再画一张偏倚风险汇总图。偏倚风险在第二个工作表里，用的 RoB 1。",
         "gold_values": g["rs-10"], "expect_code": [r"fh_forest\([^)]*rob\s*=", r"fh_rob\("], "expect_figure": True,
         "tags": ["二分类", "画图:模板", "偏倚风险"], "note": "偏倚风险判定是编的，只考流程"},
        {"id": "rs-11", "files": [f["bcg"]],
         "question": "画一张气泡图：横轴是纬度，纵轴是各研究的 RR（对数刻度），气泡大小按研究在随机效应模型里的权重。",
         "expect_figure": True, "tags": ["画图:自己画"],
         "note": "没有模板，要自己写代码画。考交付前会不会用 view_image 看一眼（报告里单独统计）"},
        {"id": "rs-12", "files": [f["nielweise_typo"]], "question": "帮我做个 meta 分析。",
         "expect_text": [r"Heard", r"超过|大于|多于|>|写反|颠倒|不可能|有误|错误|核对"], "tags": ["数据有误"],
         "note": "Heard 1998 涂层组写成了 151/5（感染数比置管数还多）。要指出来请用户核对，不能自己改数"},
    ]


# ================================================================ 标准答案
def compute_gold() -> dict[str, list]:
    from data_agent.tools.r import R_KERNEL
    from data_agent.tools.sandbox import Sandbox

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        shutil.copytree(SOURCE, work / "source")
        shutil.copy(HERE / "gold.R", work / "gold.R")
        sandbox = Sandbox.docker("finhelm-sandbox-r", R_KERNEL, work, lambda ref: {"error": "no"}, timeout_s=300)
        try:
            ex = sandbox.run('source("gold.R")')
        finally:
            sandbox.close()
    if ex.error:
        sys.exit(f"gold.R 出错：{ex.error}")
    return json.loads(re.search(r"GOLD_JSON: (.*)", ex.output).group(1))


def main() -> None:
    files = make_files()
    gold = compute_gold()
    lines = [json.dumps({"settings": {"domain": "research", "max_steps": 20}}, ensure_ascii=False)]
    lines += [json.dumps(c, ensure_ascii=False) for c in cases(files, gold)]
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"写了 {len(lines) - 1} 道题 → {OUT.relative_to(ROOT)}，{len(set(files.values()))} 个 Excel → {FILES.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
