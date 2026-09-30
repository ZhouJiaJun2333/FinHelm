"""把事件、状态、历史换成浏览器要的 JSON。

历史和事件最后都变成同一种「时间线条目」（user / assistant / tool / notice），前端只认这一种：
打开会话时用历史建出条目，之后来的事件接着改这些条目。
"""

from __future__ import annotations

import csv
import dataclasses
import datetime as dt
import re
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable

from ..core.agent import InterruptedTurn
from ..core.context import Entry, Marker
from ..core.events import Event, SubagentEvent, ToolFinished
from ..core.messages import Message
from ..core.state import AgentState
from ..tools.sandbox import Execution, saved_figures
from ..tools.sql.results import ResultStore, StoredResult

# 界面上一条工具结果最多带多少字、一张表的预览带几行
CONTENT_CHARS = 4000
PREVIEW_ROWS = 20
# 右侧面板预览 Excel / CSV 时每个工作表最多带几行
SHEET_ROWS = 500

_RESULT_REF = re.compile(r"^结果 (r\d+)（")

FileUrl = Callable[[Path], str | None]


def plain(obj: Any) -> Any:
    """递归换成 JSON 能直接写的东西。"""
    if obj is None or isinstance(obj, (str, int, float, bool)):
        return obj
    if isinstance(obj, Decimal):
        return float(obj)
    if isinstance(obj, (dt.date, dt.datetime, dt.time)):
        return obj.isoformat()
    if isinstance(obj, Path):
        return str(obj)
    if dataclasses.is_dataclass(obj) and not isinstance(obj, type):
        return {f.name: plain(getattr(obj, f.name)) for f in dataclasses.fields(obj)}
    if isinstance(obj, dict):
        return {str(k): plain(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set, frozenset)):
        return [plain(v) for v in obj]
    return str(obj)


def table_json(table: StoredResult, rows: int | None = PREVIEW_ROWS) -> dict[str, Any]:
    r = table.result
    return {"ref": table.ref, "title": table.title, "source": table.source, "sql": table.sql,
            "columns": r.columns, "rows": plain(r.rows if rows is None else r.rows[:rows]),
            "row_count": r.row_count, "truncated": r.truncated}


def details_json(details: Any, file_url: FileUrl) -> dict[str, Any] | None:
    """ToolOutput.details（只给界面的结构化结果）。认不得的就不发，界面显示模型看到的那份文字。"""
    if isinstance(details, StoredResult):
        return {"kind": "table", **table_json(details)}
    if isinstance(details, Execution):
        return {"kind": "execution", "output": details.output, "value": details.value, "error": details.error,
                "figures": [u for p in details.figures if (u := file_url(p))]}
    return None


def event_json(event: Event, file_url: FileUrl) -> dict[str, Any]:
    if isinstance(event, SubagentEvent):
        return {"type": "SubagentEvent", "data": {"task": event.task, "agent": event.agent, "title": event.title,
                                                   "event": event_json(event.event, file_url)}}
    data = {f.name: getattr(event, f.name) for f in dataclasses.fields(event)}
    if isinstance(event, ToolFinished):
        data["details"] = details_json(event.details, file_url)
        data["content"] = _cap(event.content)
    return {"type": type(event).__name__, "data": plain(data)}


def state_json(state: AgentState) -> dict[str, Any]:
    return {**plain(state), "running": state.running}


def interrupted_json(turn: InterruptedTurn | None) -> dict[str, Any] | None:
    if turn is None:
        return None
    pending = turn.pending
    return {"reason": turn.reason, "steps": turn.steps, "question": turn.question,
            "pending": None if pending is None else
            {"call_id": pending.call_id, "question": pending.question, "options": list(pending.options)}}


def timeline(entries: list[Entry] | tuple[Entry, ...], file_url: FileUrl | None = None,
             results: ResultStore | None = None) -> list[dict[str, Any]]:
    """历史 → 时间线条目。工具调用和它的结果并成一条 tool。

    历史里只有模型看到的文字，没有 details（表、图）。表按结果编号去 results 里取，图从文字里的文件清单找回来。
    """
    items: list[dict[str, Any]] = []
    tools: dict[str, dict[str, Any]] = {}
    for e in entries:
        if isinstance(e, Marker):
            items.append({"kind": "notice", "text": e.describe(), "level": "info"})
            continue
        if not isinstance(e, Message):
            continue
        if e.role == "user":
            # Agent 自己补的（推动、收尾提示、重启提醒）不是用户说的，不显示
            if not e.meta.synthetic:
                items.append({"kind": "user", "text": e.content})
        elif e.role == "assistant":
            if e.content:
                items.append({"kind": "assistant", "text": e.content, "thinking": _thinking(e.raw),
                              "streaming": False})
            for c in e.tool_calls:
                tool = {"kind": "tool", "call_id": c.id, "name": c.name, "arguments": plain(c.arguments),
                        "status": "running", "content": "", "details": None, "elapsed_ms": 0}
                tools[c.id] = tool
                items.append(tool)
        elif e.role == "tool" and (tool := tools.get(e.tool_call_id or "")) is not None:
            tool.update(status="error" if e.is_error else "done", content=_cap(e.content),
                        details=_rebuild_details(e.content, file_url, results))
    return items


def _rebuild_details(content: str, file_url: FileUrl | None, results: ResultStore | None) -> dict[str, Any] | None:
    if results is not None and (m := _RESULT_REF.match(content)) and (table := results.get(m.group(1))):
        return {"kind": "table", **table_json(table)}
    if file_url is not None and (figures := [u for p in saved_figures(content) if p.exists() and (u := file_url(p))]):
        return {"kind": "execution", "output": "", "value": None, "error": None, "figures": figures}
    return None


def sheets_json(path: Path, rows: int = SHEET_ROWS) -> list[dict[str, Any]]:
    """Excel / CSV 给右侧面板预览：每个工作表第一行当表头，其余前 rows 行。"""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
            return [_sheet(path.stem, csv.reader(f), rows)]
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [_sheet(ws.title, ws.iter_rows(values_only=True), rows) for ws in book.worksheets]
    finally:
        book.close()


def _sheet(name: str, lines: Any, rows: int) -> dict[str, Any]:
    header: list[Any] | None = None
    body: list[list[Any]] = []
    count = 0
    for line in lines:
        values = list(line)
        if header is None:
            if any(v not in (None, "") for v in values):
                header = values
            continue
        count += 1
        if len(body) < rows:
            body.append(values)
    header = header or []
    width = max([len(header), *(len(r) for r in body)], default=0)
    columns = [str(c) if c not in (None, "") else "" for c in header] + [""] * (width - len(header))
    return {"ref": name, "title": name, "source": "", "sql": "", "columns": columns,
            "rows": plain([r + [None] * (width - len(r)) for r in body]),
            "row_count": count, "truncated": count > len(body)}


def _thinking(raw: Any) -> str:
    """原生消息里的思考过程（DeepSeek 的 reasoning_content、Anthropic 的 thinking 块）。"""
    if isinstance(raw, dict):
        return raw.get("reasoning_content") or ""
    if isinstance(raw, list):
        return "\n".join(b.get("thinking", "") for b in raw if isinstance(b, dict) and b.get("type") == "thinking")
    return ""


def _cap(text: str) -> str:
    return text if len(text) <= CONTENT_CHARS else text[:CONTENT_CHARS] + f"\n…（共 {len(text)} 字）"
